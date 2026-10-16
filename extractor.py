import os
import re
import io
import openpyxl
try:
    import pandas as pd
except ImportError:
    pd = None
import pymupdf  # fitz
from PIL import Image

try:
    import easyocr
    OCR_READER = None
except ImportError:
    easyocr = None
    OCR_READER = None

def get_ocr_reader():
    global OCR_READER
    if OCR_READER is None and easyocr is not None:
        OCR_READER = easyocr.Reader(['en'], gpu=False)
    return OCR_READER

def clean_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    s = s.replace('\xa0', ' ').replace('\u2014', '-').replace('\u2013', '-')
    return s.strip()

def normalize_part_no(part_no):
    if not part_no:
        return ""
    return re.sub(r'[^A-Za-z0-9]', '', str(part_no)).upper()

def parse_price(val_str):
    if val_str is None:
        return 0.0
    if isinstance(val_str, (int, float)):
        return float(val_str)
    s = str(val_str).strip().replace(' ', '').replace('$', '').replace('RM', '').replace('MYR', '').replace('EUR', '').replace('USD', '').replace('SGD', '').replace('Euro', '')
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s and '.' not in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) <= 2:
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        matches = re.findall(r'[-+]?\d*\.?\d+', s)
        return float(matches[0]) if matches else 0.0
    except Exception:
        return 0.0

def parse_quantity(val_str):
    if val_str is None:
        return 1.0
    if isinstance(val_str, (int, float)):
        return float(val_str)
    s = str(val_str).strip()
    s = s.replace(',', '.')
    m = re.findall(r'[-+]?\d*\.?\d+', s)
    if m:
        try:
            val = float(m[0])
            return int(val) if val.is_integer() else val
        except Exception:
            return 1.0
    return 1.0

# -------------------------------------------------------------
# EXCEL PARSER (Handles RFQ and Quotations)
# -------------------------------------------------------------
def parse_excel(file_path_or_bytes):
    try:
        if isinstance(file_path_or_bytes, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
        else:
            wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)
        
        sheet = wb.active
        rows = []
        for r in sheet.iter_rows(values_only=True):
            rows.append([clean_str(c) for c in r])
    except Exception:
        # Fallback for binary .xls or unusual Excel formats using pandas
        try:
            bio = io.BytesIO(file_path_or_bytes) if isinstance(file_path_or_bytes, bytes) else file_path_or_bytes
            df = pd.read_excel(bio, header=None)
            rows = [[clean_str(c) for c in r] for r in df.values]
        except Exception:
            return []
    
    # Detect currency across sheet text
    full_sheet_text = " ".join([" ".join([str(c) for c in row if c]) for row in rows])
    detected_currency = detect_currency(full_sheet_text)

    items = []
    header_idx = -1
    col_map = {}
    
    for i, row in enumerate(rows):
        row_lower = [str(c).lower() for c in row]
        if any('item code' in c or 'part no' in c or 'article' in c or 'position' in c or 'item name' in c or 'description' in c for c in row_lower):
            header_idx = i
            for col_i, cell in enumerate(row_lower):
                # 1. Price columns FIRST to avoid collision with 'unit' in 'unit price'
                if 'unit price' in cell or 'u/price' in cell or 'u.price' in cell or 'unit_price' in cell or 'unit rate' in cell or ('price' in cell and 'total' not in cell and 'amount' not in cell) or 'rate' in cell:
                    col_map['price'] = col_i
                elif 'item code' in cell or 'item no' in cell or 'item num' in cell or 'item #' in cell:
                    col_map['item_code'] = col_i
                elif 'item name' in cell and 'parent' not in cell:
                    col_map['item_name'] = col_i
                elif 'part no' in cell or 'part number' in cell or 'p/n' in cell or 'p.n' in cell or 'p#' in cell or cell.strip() == 'pn':
                    col_map['part_no'] = col_i
                elif 'description' in cell or 'desc' in cell:
                    if 'description' not in col_map:
                        col_map['description'] = col_i
                elif 'article' in cell:
                    col_map['article_no'] = col_i
                elif 'drawing' in cell:
                    col_map['drawing_no'] = col_i
                elif 'req qty' in cell or 'rfq qty' in cell or 'qty' in cell or 'quantity' in cell:
                    if 'qty' not in col_map or 'rfq qty' in cell or 'req qty' in cell:
                        col_map['qty'] = col_i
                elif 'unit' in cell or 'uom' in cell:
                    col_map['unit'] = col_i
                elif 'other details' in cell or 'details' in cell or 'remark' in cell:
                    col_map['details'] = col_i
            break
            
    if header_idx != -1:
        current_item = None
        for i in range(header_idx + 1, len(rows)):
            row = rows[i]
            if not any(row):
                continue
            
            first_col = row[0] if len(row) > 0 else ""
            has_sno = first_col.isdigit() or (first_col and first_col.replace('.', '').isdigit())
            
            item_code_val = row[col_map.get('item_code', -1)] if 'item_code' in col_map and col_map['item_code'] < len(row) else ""
            item_name_val = row[col_map.get('item_name', -1)] if 'item_name' in col_map and col_map['item_name'] < len(row) else ""
            desc_val = row[col_map.get('description', -1)] if 'description' in col_map and col_map['description'] < len(row) else ""
            part_no_val = row[col_map.get('part_no', -1)] if 'part_no' in col_map and col_map['part_no'] < len(row) else ""
            qty_val = row[col_map.get('qty', -1)] if 'qty' in col_map and col_map['qty'] < len(row) else ""
            unit_val = row[col_map.get('unit', -1)] if 'unit' in col_map and col_map['unit'] < len(row) else ""
            price_val = row[col_map.get('price', -1)] if 'price' in col_map and col_map['price'] < len(row) else ""
            
            full_row_text = " ".join(row)
            
            if has_sno or (item_code_val and not current_item) or (item_name_val and qty_val):
                if current_item:
                    items.append(current_item)
                
                extracted_part = part_no_val
                if not extracted_part or extracted_part.lower() in ['n/a', 'na', '-']:
                    m_part = re.search(r'\b(?:Part\s*(?:Number|No|#)|P/N)\s*[:=\-]?\s*([A-Za-z0-9\-\./]+)', full_row_text, re.IGNORECASE)
                    if m_part:
                        extracted_part = m_part.group(1).strip()
                
                final_code = item_code_val or extracted_part or (row[col_map['article_no']] if 'article_no' in col_map and col_map['article_no'] < len(row) else "")
                final_desc = item_name_val or desc_val or ""
                
                current_item = {
                    'item_no': len(items) + 1,
                    'item_code': final_code,
                    'part_no': extracted_part or "",
                    'description': final_desc,
                    'quantity': parse_quantity(qty_val),
                    'unit': unit_val or 'PCS',
                    'unit_price': parse_price(price_val),
                    'currency': detected_currency,
                    'raw_details': full_row_text
                }
            else:
                if current_item:
                    m_part = re.search(r'\b(?:Part\s*(?:Number|No|#)|P/N)\s*[:=\-]?\s*([A-Za-z0-9\-\./]+)', full_row_text, re.IGNORECASE)
                    if m_part:
                        part_extracted = m_part.group(1).strip()
                        if len(part_extracted) >= 2 and not part_extracted.lower() in ['parts', 'shipping']:
                            current_item['part_no'] = part_extracted
                            if not current_item['item_code']:
                                current_item['item_code'] = part_extracted
                    current_item['raw_details'] += " | " + full_row_text
                    
        if current_item:
            items.append(current_item)
            
    return items

# -------------------------------------------------------------
# PDF PARSER (Vector & Digital PDFs)
# -------------------------------------------------------------
def parse_pdf(file_path_or_bytes):
    if isinstance(file_path_or_bytes, bytes):
        doc = pymupdf.open(stream=file_path_or_bytes, filetype="pdf")
    else:
        doc = pymupdf.open(file_path_or_bytes)
        
    full_text = ""
    for page in doc:
        full_text += page.get_text() + "\n"
        
    if len(full_text.strip()) < 50:
        return parse_scanned_pdf(doc)
    
    currency = detect_currency(full_text)
    items = []
    
    # Marine Requisition Table Format (e.g. Dinastia Jati / DJ Marine Services style)
    # Tables with columns: SNo, Item Code (e.g. 11.02.01.26), Item Name, Other Details (Make, Part Number, ESN)
    req_matches = list(re.finditer(r'(?:^|\n)(\d+)\s*\n(?:([^\n]+)\s*\n)?([0-9]{2}\.[0-9]{2}\.[0-9]{2}\.[0-9]{2})\s*\n(.*?)(?=\n\d+\s*\n(?:[^\n]+\s*\n)?[0-9]{2}\.[0-9]{2}|\Z)', full_text, re.DOTALL))
    if len(req_matches) >= 1:
        for m in req_matches:
            sno = m.group(1)
            parent_name = m.group(2) or ''
            code = m.group(3)
            chunk = m.group(4)

            # Part Number extraction
            m_p = re.search(r'(?:Part\s*(?:Number|No|#)|P/N)\s*[:=\s]*([A-Za-z0-9\-\./]+)', chunk, re.IGNORECASE)
            part_no = m_p.group(1).strip() if m_p else ''

            # Model / Make extraction
            m_maker = re.search(r'(?:Make|Model(?:\s*Number)?)\s*[:=\s]*([^\n,;]+)', chunk, re.IGNORECASE)
            maker_no = m_maker.group(1).strip() if m_maker else ''

            # Product description lines (excluding metadata words and dates)
            desc_lines = [l.strip() for l in chunk.split('\n') if l.strip() and not l.strip().isdigit() and not any(k in l for k in ['Other Details', 'Comments:', 'Make:', 'Model Number:', 'Part Number', 'ESN', 'Drawing Number']) and not re.match(r'^\d{2}/\d{2}/\d{2,4}$', l.strip()) and l.strip() not in ['Nos', 'Pcs', 'No', 'Yes', '0.00', '1.00']]
            desc = ' '.join(desc_lines[:2]).strip()
            if parent_name and not desc:
                desc = parent_name.strip()
            elif parent_name and parent_name.strip().lower() not in desc.lower() and len(parent_name.strip()) <= 40:
                desc = f"{parent_name.strip()} - {desc}".strip()

            # Quantity and Unit (strip dates like 04/10/202\n4 or 04/10/2024 first)
            qty = 1.0
            unit = 'PCS'
            clean_chunk = re.sub(r'\d{2}/\d{2}/\d{2,4}(?:\s*\n\s*\d+)?', ' ', chunk)
            m_qty = re.search(r'\b(\d+(?:\.\d+)?)\s*\n\s*(?:(\d+(?:\.\d+)?)\s*\n\s*)?(Pcs|Nos|Set|Sets|Pce|Pces|Box|Ea|Each|Mtr)\b', clean_chunk, re.IGNORECASE)
            if m_qty:
                qty = float(m_qty.group(1))
                unit = (m_qty.group(3) or m_qty.group(2) or 'PCS').upper()

            items.append({
                'item_no': int(sno),
                'item_code': code,
                'part_no': part_no,
                'maker_no': maker_no,
                'description': desc or f"Item {sno}",
                'quantity': qty,
                'unit': unit,
                'unit_price': 0.0,
                'currency': currency,
                'raw_details': chunk[:200]
            })

        if items:
            return items, currency

    for page in doc:
        blocks = page.get_text('blocks')
        for b in blocks:
            b_text = clean_str(b[4])
            if not b_text:
                continue
            b_lines = [clean_str(l) for l in b_text.split('\n') if clean_str(l)]
            
            # Skip non-table info blocks
            if any(h in b_text.lower() for h in ['payment terms', 'validity', 'beneficiary', 'swift', 'www.', 'tel:', 'fax:', 'bank', 'ex-works', 'total value', 'subtotal', 'comments:']):
                continue
                
            # Block Format 1: Oceantech style
            if len(b_lines) >= 5 and re.match(r'^\d+[.,]\d{2}$', b_lines[-1]) and re.match(r'^\d+[.,]\d{2}$', b_lines[-2]):
                desc = b_lines[0]
                part_no = b_lines[1]
                qty = parse_quantity(b_lines[2])
                unit = b_lines[3]
                price = parse_price(b_lines[-2])
                
                items.append({
                    'item_no': len(items) + 1,
                    'item_code': part_no,
                    'part_no': part_no,
                    'description': desc,
                    'quantity': qty,
                    'unit': unit,
                    'unit_price': price,
                    'currency': currency
                })
                continue
                
            # Block Format 2: Challenge Power & Diomar style
            if len(b_lines) >= 6 and b_lines[0].isdigit() and re.search(r'\d+[.,]\d{2}', b_lines[-1]):
                num = int(b_lines[0])
                desc = b_lines[1]
                part_no = ""
                
                if b_lines[1].isdigit() and len(b_lines[1]) >= 3:
                    part_no = b_lines[1]
                    desc = b_lines[2]
                else:
                    m_pn = re.search(r'(?:P/N|PN|CODE)\s*[:#]?\s*([A-Za-z0-9\-_]+)', desc, re.IGNORECASE)
                    if m_pn:
                        part_no = m_pn.group(1)
                    else:
                        m_rep = re.search(r'\b(\d{4}[A-Za-z0-9\-]*)\b', desc)
                        part_no = m_rep.group(1) if m_rep else str(num)
                        
                price_matches = [parse_price(x) for x in b_lines if re.search(r'\d+[.,]\d{2}', x)]
                unit_price = price_matches[0] if price_matches else 0.0
                
                unit = "PCS"
                qty = 1.0
                unit_idx = next((i for i, x in enumerate(b_lines) if x.upper() in ['PCS', 'PC', 'PCE', 'SET', 'ST', 'UNIT', 'EA']), -1)
                if unit_idx > 0:
                    unit = b_lines[unit_idx].upper()
                    if b_lines[unit_idx - 1].isdigit():
                        qty = parse_quantity(b_lines[unit_idx - 1])
                else:
                    for x in b_lines[2:]:
                        if x.isdigit() and int(x) != num and int(x) != int(part_no if part_no.isdigit() else 0) and parse_price(x) not in price_matches:
                            qty = parse_quantity(x)
                            break
                
                items.append({
                    'item_no': num,
                    'item_code': part_no or str(num),
                    'part_no': part_no,
                    'description': desc,
                    'quantity': qty,
                    'unit': unit,
                    'unit_price': unit_price,
                    'currency': currency
                })
                continue

    if items:
        return items, currency
        
    # Aflex / Sequential Line parser fallback
    lines = [clean_str(l) for l in full_text.split('\n') if clean_str(l)]
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.isdigit() and int(line) == len(items) + 1 and i + 4 < len(lines):
            item_no = int(line)
            desc_line = lines[i+1]
            qty_line = lines[i+2]
            uom_line = lines[i+3]
            price_line = lines[i+4]
            
            if re.match(r'^\d+(?:\.\d+)?$', qty_line.replace(',', '.')) and (re.search(r'\d+[.,]\d{2}', price_line) or price_line.replace(',', '').replace('.', '').isdigit()):
                part_no = ""
                m_paren = re.search(r'\(([A-Za-z0-9\-_]+)\)', desc_line)
                m_pn = re.search(r'(?:PN|P/N|Part\s*No)?[:\s]*([A-Za-z0-9\-_]+)', desc_line, re.IGNORECASE)
                if m_paren:
                    part_no = m_paren.group(1)
                elif m_pn:
                    part_no = m_pn.group(1)
                
                items.append({
                    'item_no': item_no,
                    'item_code': part_no or str(item_no),
                    'part_no': part_no,
                    'description': desc_line,
                    'quantity': parse_quantity(qty_line),
                    'unit': uom_line,
                    'unit_price': parse_price(price_line),
                    'currency': currency
                })
                i += 5
                continue
        i += 1

    return items, currency

# -------------------------------------------------------------
# SCANNED PDF & IMAGE / SCREENSHOT OCR PARSER
# -------------------------------------------------------------
def parse_scanned_pdf(doc):
    images = []
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        images.append(img)
    return parse_images_ocr(images)

def parse_image_file(image_path_or_bytes):
    if isinstance(image_path_or_bytes, bytes):
        img = Image.open(io.BytesIO(image_path_or_bytes))
    else:
        img = Image.open(image_path_or_bytes)
    return parse_images_ocr([img])

def parse_images_ocr(images):
    reader = get_ocr_reader()
    if reader is None:
        return [], "MYR"
        
    all_items = []
    detected_currency = "MYR"
    
    for img in images:
        img_byte_arr = io.BytesIO()
        img.save(img_byte_arr, format='PNG')
        res = reader.readtext(img_byte_arr.getvalue())
        
        full_text = " ".join([b[1] for b in res])
        curr = detect_currency(full_text)
        if curr != "MYR":
            detected_currency = curr
            
        page_items = parse_single_image_ocr(res, detected_currency)
        all_items.extend(page_items)
        
    for idx, it in enumerate(all_items, start=1):
        it['item_no'] = idx
        
    return all_items, detected_currency

def parse_single_image_ocr(ocr_results, currency="MYR"):
    """
    Intelligently identifies table header and footer boundaries,
    filters out company info, contacts, notes, and extracts clean item rows.
    """
    header_boxes = []
    for b in ocr_results:
        t = b[1].strip().upper()
        if any(t == k or t.startswith(k + ' ') or ('ITEM CODE' in t and '*' in t) for k in ['NO.', 'NO', 'ITEM NO', 'ITEM CODE', 'ITEM NAME', 'DESCRIPTION', 'CODE', 'P/N', 'QTY', 'Q\'TY', 'QUANTITY', 'REQ QTY', 'RFQ QTY', 'UPRICE', 'U/PRICE', 'PRICE', 'AMOUNT', 'SNO', 'OTHER DETAILS']):
            if b[0][0][1] < 1200:
                header_boxes.append(b)
                
    header_bottom = max(b[0][2][1] for b in header_boxes) if header_boxes else 0
    
    col_desc_x = next((b[0][0][0] for b in header_boxes if 'DESC' in b[1].upper() or 'ITEM NAME' in b[1].upper()), 100)
    col_code_x = next((b[0][0][0] for b in header_boxes if 'CODE' in b[1].upper() or 'P/N' in b[1].upper() or 'PART' in b[1].upper()), 650)
    col_qty_x = next((b[0][0][0] for b in header_boxes if 'QTY' in b[1].upper() or 'Q\'TY' in b[1].upper()), 950)
    col_price_x = next((b[0][0][0] for b in header_boxes if 'PRICE' in b[1].upper() or 'UPRICE' in b[1].upper() or 'U/PRICE' in b[1].upper()), 1200)
    
    footer_y = 99999
    for b in ocr_results:
        t = b[1].upper()
        if any(k in t for k in ['TOTAL=', 'TOTAL =', '<REMARK>', 'TO BE CONTINUED', 'MADE IN KOREA', 'AFTER ORDER ACK', 'TOTAL EX-WORK', 'BENEFICIARY:', 'COMMENTS:']):
            if b[0][0][1] > header_bottom and b[0][0][1] < footer_y:
                footer_y = b[0][0][1]
                
    body_boxes = [b for b in ocr_results if b[0][0][1] > header_bottom and b[0][2][1] < footer_y]
    body_boxes = [b for b in body_boxes if not any(k in b[1].upper() for k in ['MATERIAL:', 'SYMBOL:', 'CAC502', 'CAC402', 'SUS304', 'SCM435', 'SS400', 'CR-MO', 'DIMENSION', 'WEIGHT =', 'TYPE APPROVED', 'DRAING NO:', 'MAIN ENGIN', 'WATER PUMP'])]
    
    # 1. Multi-line Block RFQ Parsing (e.g. sample.png / sample2.png)
    anchors = []
    for b in body_boxes:
        txt = b[1].strip()
        x = b[0][0][0]
        y = (b[0][0][1] + b[0][2][1]) / 2.0
        if re.search(r'M\d{3}\.\d{3}/|\d{3}\.\d{3}/\d{5}|Our\s*Part\s*No', txt, re.IGNORECASE) or (x < 100 and re.match(r'^[1-9]\d?$', txt)):
            anchors.append(y)
            
    if len(anchors) >= 1 and any(re.search(r'M\d{3}\.\d{3}/|\d{3}\.\d{3}/|Part\s*Number|Range|Pressure\s*Gauge', b[1], re.IGNORECASE) for b in body_boxes):
        anchors.sort()
        merged_anchors = []
        for a in anchors:
            if not merged_anchors or abs(a - merged_anchors[-1]) > 45.0:
                merged_anchors.append(a)
                
        clusters = [[] for _ in merged_anchors]
        for b in body_boxes:
            y = (b[0][0][1] + b[0][2][1]) / 2.0
            closest_idx = min(range(len(merged_anchors)), key=lambda i: abs(y - merged_anchors[i]))
            clusters[closest_idx].append(b)
            
        items = []
        for c in clusters:
            if not c:
                continue
            c.sort(key=lambda b: (b[0][0][1], b[0][0][0]))
            cluster_text = " ".join([b[1] for b in c])
            
            # Item Code
            m_code = re.search(r'\b(M\d{3}\.\d{3}/[0-9A-Za-z]+|\d{3}\.\d{3}/\d{5})\b', cluster_text)
            item_code = m_code.group(1).replace('I', '1').replace('O', '0') if m_code else ''
            
            # Part Number
            m_part = re.search(r'(?:Part\s*(?:Number|No)?|Our\s*Part\s*No[:=\s]*)\s*([0-9A-Za-z\-_\./]+)', cluster_text, re.IGNORECASE)
            part_no = m_part.group(1).replace('IW-', '1W-') if m_part else item_code
            if not item_code and part_no:
                item_code = part_no
                
            # Quantity
            qty = 1.0
            m_qty = re.search(r'\b(\d+(?:[.,]\d+)?)\s*(?:PC|PCS|SET|EA|ST|UNIT)\b', cluster_text, re.IGNORECASE)
            if m_qty:
                qty = float(m_qty.group(1).replace(',', '.'))
            else:
                qtys = [float(b[1].replace(',', '.')) for b in c if re.match(r'^\d+(?:[.,]\d+)?$', b[1]) and float(b[1].replace(',', '.')) < 500]
                if qtys:
                    qty = qtys[0]
                    
            # Unit
            m_u = re.search(r'\b(PCS|PC|PCE|SET|ST|UNIT|EA)\b', cluster_text, re.IGNORECASE)
            unit = m_u.group(1).upper() if m_u else 'PCS'
            
            # Price
            prices = [parse_price(b[1]) for b in c if re.match(r'^\d+[.,]\d{2}$', b[1])]
            unit_price = prices[-1] if prices else 0.0
            
            # Description
            desc = cluster_text
            m_caterpillar = re.search(r'([FR]\.O\s+HIGH\s+PRESSURE\s+PIPE\s+UNIT\s+NO\.\d+)', cluster_text, re.IGNORECASE)
            if m_caterpillar:
                desc = m_caterpillar.group(1).replace('R.O', 'F.O')
            else:
                desc = re.sub(r'Make\s*[:=]?\s*[^|\n]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'Model Number\s*[:=]?\s*[^|\n]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'Part Number\s*[:=]?\s*[^|\n]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'Our Part No\s*[:=]?\s*[\d\./]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'Drawing Number\s*[:=]?\s*[^|\n]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'Parent Item Name\s*[:=]?\s*[^|\n]+', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'FUEL SYSTEM', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'SHIP EQUIPMENT', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'\b(M\d{3}\.\d{3}/[0-9A-Za-z]+|\d{3}\.\d{3}/\d{5})\b', '', desc)
                desc = re.sub(r'\b(PCS|PC|PCE|SET|ST|UNIT|EA|NA|TBD)\b', '', desc, flags=re.IGNORECASE)
                desc = re.sub(r'\b\d+(?:[.,]\d+)?\b', '', desc)
                desc = re.sub(r'\s+', ' ', desc).strip()
            
            if item_code or desc:
                items.append({
                    'item_code': item_code or part_no,
                    'part_no': part_no or item_code,
                    'description': desc or item_code,
                    'quantity': int(qty) if qty.is_integer() else qty,
                    'unit': unit,
                    'unit_price': unit_price
                })
        return items

    # 2. Standard Row-by-Row Quotation Extraction (M260378 / M260379 / etc.)
    body_boxes.sort(key=lambda b: (b[0][0][1], b[0][0][0]))
    rows = []
    curr = []
    curr_y = None
    for b in body_boxes:
        y = (b[0][0][1] + b[0][2][1]) / 2.0
        if curr_y is None:
            curr_y = y
            curr.append(b)
        elif abs(y - curr_y) <= 24.0:
            curr.append(b)
            curr_y = sum((x[0][0][1] + x[0][2][1]) / 2.0 for x in curr) / len(curr)
        else:
            rows.append(curr)
            curr = [b]
            curr_y = y
    if curr:
        rows.append(curr)
        
    items = []
    for r in rows:
        desc_tokens = []
        code_tokens = []
        qty_tokens = []
        price_tokens = []
        
        for b in r:
            x = (b[0][0][0] + b[0][1][0]) / 2.0
            txt = b[1].strip()
            if x < col_code_x - 30:
                if x < col_desc_x - 20 and (txt.isdigit() or txt.replace('.', '').isdigit()):
                    pass
                else:
                    desc_tokens.append(txt)
            elif x < col_qty_x - 30:
                code_tokens.append(txt)
            elif x < col_price_x - 20:
                qty_tokens.append(txt)
            else:
                price_tokens.append(txt)
                
        desc = ' '.join(desc_tokens).strip()
        code = ' '.join(code_tokens).strip()
        qty_str = ' '.join(qty_tokens).strip()
        price_str = ' '.join(price_tokens).strip()
        
        code = code.replace('OZ0', '020').replace('O20', '020').replace('O21', '021')
        
        p_matches = re.findall(r'\b\d+[.,]\d{2}\b', price_str)
        if not p_matches:
            p_matches = re.findall(r'\b\d+[.,]\d{2}\b', ' '.join([b[1] for b in r]))
            
        if not p_matches or not desc:
            continue
            
        u_price = float(p_matches[0].replace(' ', '').replace(',', '.'))
        if u_price <= 0:
            continue
            
        m_q = re.search(r'(\d+)\s*(PCS|SET|PC|PCE|EA|ST)?', qty_str, re.IGNORECASE)
        qty = int(m_q.group(1)) if m_q else 1
        unit = m_q.group(2).upper() if m_q and m_q.group(2) else 'PCS'
        
        items.append({
            'item_code': code or desc,
            'part_no': code or '',
            'description': desc or code,
            'quantity': qty,
            'unit': unit,
            'unit_price': u_price
        })
        
    return items

# -------------------------------------------------------------
# -------------------------------------------------------------
# RAW TEXT / EMAIL BODY PARSER
# -------------------------------------------------------------
NOISE_TEXT_PATTERNS = [
    'TEL:', 'FAX:', 'EMAIL:', 'WWW.', 'HTTP:', 'HTTPS:', 'BEST REGARDS', 'KIND REGARDS', 
    'REGARDS', 'DEAR SIR', 'DEAR MADAM', 'DEAR SALES', 'DEAR ', 'GOOD DAY',
    'QUOTATION NO:', 'TOTAL VALUE', 'BANK', 'DISCLAIMER', 'CONFIDENTIAL', 'VIRUS',
    'IMMUNITY', 'PRIVILEGED', 'ADDRESSEE', 'VAKCHOU', 'TOH GUAN', 'ENTERPRISE HUB',
    'HIGH STREET', 'SINGAPORE', 'GREECE', 'KOREA', 'BUSAN', 'NICOSIA', 'CYPRUS',
    'EX-WORKS', 'PAYMENT TERMS', 'VALIDITY', 'DELIVERY TIME', 'SUBJECT TO PRIOR',
    'PACKING AND DELIVERY', 'PACKING CHARGE', 'THANK YOU', 'WE ARE PLEASED',
    'RECEIPT CONFIRMATION', 'REPLY TO ALL', 'PROCUREMENT EXECUTIVE', 'WEBSITE:',
    'OUT OF OUR PRODUCT RANGE', 'UNABLE TO QUOTE', 'REGRET TO INFORM', 'DECLINE TO QUOTE',
    'PLEASE ADVISE', 'PLEASE PROVIDE', 'OUTLINE DRAWING', 'TECHNICAL DATA SHEET',
    'SPARE PARTS LIST', 'AFTER RECEIPT OF ORDER', 'WORKING DAYS', 'AP MERIN'
]

def parse_raw_text(text_content):
    text_content = clean_str(text_content)
    if not text_content:
        return [], "MYR"
        
    lines = [l.strip() for l in text_content.split('\n') if l.strip()]
    currency = detect_currency(text_content)
    items = []
    
    for i, line in enumerate(lines):
        line_upper = line.upper()
        
        # Skip noise, greetings, disclaimers, contact details, addresses
        if any(pat in line_upper for pat in NOISE_TEXT_PATTERNS):
            continue
            
        # Delimited lines (tabs or pipes)
        if '\t' in line or '|' in line:
            delim = '\t' if '\t' in line else '|'
            parts = [clean_str(p) for p in line.split(delim) if clean_str(p)]
            if len(parts) >= 3:
                code = parts[0]
                desc = parts[1]
                qty = parse_quantity(parts[2]) if len(parts) > 2 else 1
                unit = parts[3] if len(parts) > 3 and not parts[3].replace('.', '').isdigit() else "PCS"
                price = parse_price(parts[-1]) if len(parts) >= 4 else 0.0
                
                # Only accept valid commercial items with price > 0
                if price > 0 and len(desc) >= 3 and not any(pat in desc.upper() for pat in NOISE_TEXT_PATTERNS):
                    items.append({
                        'item_no': len(items) + 1,
                        'item_code': code,
                        'part_no': code,
                        'description': desc,
                        'quantity': qty,
                        'unit': unit,
                        'unit_price': price,
                        'currency': currency
                    })
        elif ',' in line:
            # Comma-separated: ONLY if last element is a valid price > 0
            parts = [clean_str(p) for p in line.split(',') if clean_str(p)]
            if len(parts) >= 3:
                last_token = parts[-1].strip()
                if re.match(r'^[A-Z]{0,3}\s*\$?\s*\d+(?:[.,]\d{2})?$', last_token, re.IGNORECASE):
                    price = parse_price(last_token)
                    if price > 0:
                        code = parts[0]
                        desc = parts[1]
                        qty = parse_quantity(parts[2]) if len(parts) > 3 else 1
                        unit = parts[3] if len(parts) > 4 and not parts[3].replace('.', '').isdigit() else "PCS"
                        if len(desc) >= 3 and not any(pat in desc.upper() for pat in NOISE_TEXT_PATTERNS):
                            items.append({
                                'item_no': len(items) + 1,
                                'item_code': code,
                                'part_no': code,
                                'description': desc,
                                'quantity': qty,
                                'unit': unit,
                                'unit_price': price,
                                'currency': currency
                            })
        else:
            m = re.match(r'^(\d+)?[\.\s]*([A-Za-z0-9\-\./]+)?\s+(.+?)\s+(\d+(?:[.,]\d+)?)\s*([A-Za-z]+)?\s+([0-9\s.,]+)$', line)
            if m:
                u_p = parse_price(m.group(6)) if m.group(6) else 0.0
                desc = m.group(3).strip()
                if u_p > 0 and len(desc) >= 3 and not any(pat in desc.upper() for pat in NOISE_TEXT_PATTERNS):
                    items.append({
                        'item_no': len(items) + 1,
                        'item_code': m.group(2) or str(len(items) + 1),
                        'part_no': m.group(2) or "",
                        'description': desc,
                        'quantity': parse_quantity(m.group(4)),
                        'unit': m.group(5) or "PCS",
                        'unit_price': u_p,
                        'currency': currency
                    })
                
    return items, currency

def detect_currency(text):
    t = text.upper()
    if 'USD' in t or 'US DOLLAR' in t or 'USS' in t or 'US$' in t:
        return 'USD'
    elif 'EUR' in t or 'EURO' in t or '€' in t:
        return 'EUR'
    elif 'SGD' in t or 'SINGAPORE DOLLAR' in t or 'S$' in t:
        return 'SGD'
    elif 'GBP' in t or 'POUND' in t or '£' in t:
        return 'GBP'
    elif 'JPY' in t or 'YEN' in t or '¥' in t:
        return 'JPY'
    elif 'CNY' in t or 'RMB' in t:
        return 'CNY'
    elif 'MYR' in t or 'RINGGIT' in t or 'RM' in t:
        return 'MYR'
    return 'MYR'
