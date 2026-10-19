"""Incremental IMAP ingestion. Mail is evidence, never executable instructions."""
import email
import hashlib
import imaplib
import io
import json
import os
import re
import ssl
import uuid
from email import policy
from email.utils import getaddresses
from pathlib import Path
from .store import now

def clean_body(text):
    text=re.split(r'(?im)^On .+wrote:|^-{2,}\s*Original Message|^From:\s|^Sent:\s',text)[0]
    text=re.split(r'(?m)^--\s*$',text)[0]
    return text.strip()[:16000]

def parse_message(raw):
    msg=email.message_from_bytes(raw,policy=policy.default)
    text=[];html=[];attachments=[]
    for part in msg.walk():
        if part.is_multipart():continue
        payload=part.get_payload(decode=True) or b''
        filename=part.get_filename() or ('inline-image' if part.get_content_type().startswith('image/') and part.get('Content-ID') else None)
        if filename:
            if len(payload)>20*1024*1024:raise ValueError('Attachment exceeds 20 MB review limit')
            attachments.append(dict(filename=filename,hash=hashlib.sha256(payload).hexdigest(),bytes=payload,content_type=part.get_content_type(),content_id=str(part.get('Content-ID','')).strip('<>'),disposition=part.get_content_disposition()))
        elif part.get_content_type()=='text/plain':
            text.append(payload.decode(part.get_content_charset() or 'utf-8',errors='replace'))
        elif part.get_content_type()=='text/html':html.append(payload.decode(part.get_content_charset() or 'utf-8',errors='replace'))
    if not text and html:
        try:
            from bs4 import BeautifulSoup
            text=[BeautifulSoup('\n'.join(html),'html.parser').get_text('\n')]
        except ImportError:text=[re.sub('<[^>]+>',' ','\n'.join(html))]
    identity_body=clean_body('\n'.join(text))
    if html and not '\n'.join(text).strip():
        from bs4 import BeautifulSoup
        text=[BeautifulSoup('\n'.join(html),'html.parser').get_text('\n')]
    sender=getaddresses([str(msg.get('From',''))])
    return dict(message_id=str(msg.get('Message-ID','')).strip(),references=re.findall(r'<[^>]+>',str(msg.get('References',''))+' '+str(msg.get('In-Reply-To',''))),subject=str(msg.get('Subject','')),sender=(sender[0][1] if sender else '').lower(),recipients=[a.lower() for _,a in getaddresses([str(msg.get('To','')),str(msg.get('Cc',''))])],date=str(msg.get('Date','')),body=clean_body('\n'.join(text)),identity_body=identity_body,body_full='\n'.join(text),body_html='\n'.join(html),from_header=str(msg.get('From','')),to_header=str(msg.get('To','')),cc_header=str(msg.get('Cc','')),attachments=attachments)

def classify(data):
    text=(data['subject']+'\n'+data['body']).lower()
    # A new client invitation can mention quotations, shipping and technical details.
    # Replies still follow the event checks below.
    if re.match(r'^\s*request for quot(?:e|ation)\b',data['subject'],re.I):return 'rfq'
    if re.search(r'\b(?:please|kindly|could you|can you)\b.{0,100}\b(?:vessel|datasheet|data sheet|serial|confirm|clarif|details|drawing)',text,re.S):return 'clarification'
    if re.search(r'\b(?:airway bill|awb|bill of lading|delivery order|tracking number)\b',text):return 'shipping_evidence'
    if re.search(r'\b(?:purchase order|po attached|order confirmation)\b',text):return 'order_evidence'
    if re.search(r'\b(?:remittance|payment advice|paid|payment made)\b',text):return 'payment_evidence'
    if re.search(r'\b(?:please|kindly)\s+(?:quote\b|(?:provide|submit|send)\b.{0,35}\b(?:quotation|quote|offer)\b)|\brequest for quotation\b|\bwe need (?:a |your )?(?:quotation|quote)\b',text):return 'rfq'
    if re.match(r'^\s*(?:rfq\b|request for quotation\b)',data['subject'],re.I):return 'rfq'
    if re.search(r'\b(?:quotation|our quote|offer attached)\b',text):return 'supplier_quote'
    if re.search(r'\b(?:rfq|request for quotation|requisition)\b',text):return 'rfq'
    return 'unknown'

def candidate_jobs(store,c,data):
    from .rfq_identity import facts,normalize
    ids=set()
    for ref in data.get('references',[]):
        for r in c.execute('SELECT job_id FROM messages WHERE message_id=? AND job_id IS NOT NULL',(ref,)):ids.add(r[0])
    text=data['subject']+' '+data['body']
    explicit=re.findall(r'\bJC-\d{2}-[A-Z0-9]+\b',text,re.I)
    for jid in explicit:
        if c.execute('SELECT 1 FROM jobs WHERE id=?',(jid.upper(),)).fetchone():ids.add(jid.upper())
    identity=facts(data)
    for row in c.execute('SELECT id FROM jobs'):
        job=store.job(c,row[0])
        if identity['vessel_name'] and normalize(job['vessel_name'])!=identity['vessel_name']:continue
        for req in [job['requisition_no'],job.get('client_reference',''),job.get('company_reference','')]:
            if req and re.search(r'(?<![\w])'+re.escape(req)+r'(?![\w])',text,re.I):ids.add(job['id'])
    return sorted(ids)

def known_party(store,c,job,data):
    suppliers={r.get('supplier_email','').lower() for r in store.records(c,job['id']) if r['kind'] in ['supplier_rfq','supplier_po','supplier_quote']}
    return data['sender']==job['client_email'].lower() or data['sender'] in suppliers

def apply_event(store,flow,c,jid,mid,data,event,actor):
    job=store.job(c,jid)
    if event=='supplier_quote' and data.get('sender','').lower()==job.get('client_email','').lower() and classify(data)=='rfq':event='rfq'
    evidence='email:'+mid
    if event=='clarification':
        flow.record(c,jid,'clarification',dict(version=job['version'],question=data['body'] or data['subject'],supplier_email=data['sender'],evidence=evidence),actor)
    elif event=='supplier_quote':
        # Append supplier evidence; never roll an awarded job back to quotes_received.
        store.add_record(c,jid,'supplier_quote',dict(supplier_email=data['sender'],evidence=evidence,subject=data['subject']),actor)
        store.bump(c,jid)
    else:
        store.event(c,jid,actor,event,dict(evidence=evidence,subject=data['subject'],body=data['body']))
        store.bump(c,jid)

def ingest(store,flow,raw,mailbox,folder,validity,uid,live_monitor=True,mailbox_owner=''):
    data=parse_message(raw)
    # Include payload hashes to distinguish revised contents even if Message-ID was reused.
    # Rich presentation fields must not change the existing ingestion identity.
    content=json.dumps({k:(data.get('identity_body',data[k]) if k=='body' else data[k]) for k in ['message_id','references','subject','sender','date','body']},sort_keys=True)+''.join(a['hash'] for a in data['attachments'])
    digest=hashlib.sha256(content.encode()).hexdigest()
    mid=hashlib.sha256((data['message_id']+'|'+digest).encode()).hexdigest()
    raw_path=store.blob('original_emails',mid+'.eml')
    if not raw_path.exists():raw_path.write_bytes(raw)
    for a in data['attachments']:
        path=store.blob('attachments',a['hash'])
        if not path.exists():path.write_bytes(a['bytes'])
        del a['bytes']
    with store.transaction() as c:
        if c.execute('SELECT 1 FROM occurrences WHERE mailbox=? AND folder=? AND validity=? AND uid=?',(mailbox,folder,validity,uid)).fetchone():return 'duplicate'
        if c.execute('SELECT 1 FROM messages WHERE id=?',(mid,)).fetchone():
            c.execute('INSERT INTO occurrences VALUES (?,?,?,?,?)',(mailbox,folder,validity,uid,mid));return 'duplicate'
        candidates=candidate_jobs(store,c,data)
        event=classify(data)
        data.update(candidates=candidates,event=event,direction='sent' if folder.lower()!='inbox' else 'received')
        data['direction']='sent' if 'sent' in folder.lower() or data['sender']==str(mailbox).lower() or data['sender'].endswith('@apmerinservices.com.my') else 'received'
        data['new_rfq']=not candidates and event=='rfq' and data['direction']=='received'
        from .rfq_identity import facts
        data.update(live_monitor=live_monitor,mailbox_owner=mailbox_owner,rfq_suggestion=facts(data))
        jid=candidates[0] if len(candidates)==1 else None
        status='review'
        auto=store.setting('automation_mode','review',c)=='automatic'
        if jid and data['direction']=='received' and known_party(store,c,store.job(c,jid),data) and event in ['clarification','supplier_quote']:
            job=store.job(c,jid)
            is_supplier=any(r.get('supplier_email','').lower()==data['sender'] for r in store.records(c,jid) if r['kind'] in ['supplier_rfq','supplier_po'])
            if auto and is_supplier:
                status='applied'
        # Incoming RFQs alert first. Staff sent acknowledgements / sourcing start the card.
        c.execute('INSERT INTO messages VALUES (?,?,?,?,?,?,?)',(mid,data['message_id'],digest,json.dumps(data),jid,status,now()))
        c.execute('INSERT INTO occurrences VALUES (?,?,?,?,?)',(mailbox,folder,validity,uid,mid))
        if status=='applied':apply_event(store,flow,c,jid,mid,data,event,'worker')
        elif jid:store.event(c,jid,'worker','email_review',dict(evidence='email:'+mid,event=event))
    return status

def approve_message(store,flow,c,mid,decision,actor):
    row=c.execute('SELECT * FROM messages WHERE id=?',(mid,)).fetchone()
    if not row:raise LookupError('Message not found')
    if row['status'] not in ['review','ai_pending']:raise RuntimeError('Message already reviewed')
    data=json.loads(row['data'])
    if decision.get('action')=='ignore':
        c.execute("UPDATE messages SET status='ignored' WHERE id=?",(mid,));return 'ignored'
    jid=decision.get('job_id')
    if decision.get('action')=='create':
        from .rfq_identity import same_job
        identity=decision.get('job',{})
        matches=[store.job(c,r[0]) for r in c.execute('SELECT id FROM jobs') if same_job(store.job(c,r[0]),identity)]
        if len(matches)>1:raise ValueError('Several jobs match this RFQ. Select the correct existing job.')
        owner=actor if c.execute("SELECT 1 FROM users WHERE id=? AND role='procurement' AND active=1",(actor,)).fetchone() else ''
        job=matches[0] if matches else flow.create(c,dict(identity,owner=identity.get('owner') or owner),actor);jid=job['id']
    if not jid:raise ValueError('Select a job or create one')
    store.job(c,jid)
    event=decision.get('event',data['event'])
    if event not in ['clarification','supplier_quote','rfq','shipping_evidence','order_evidence','payment_evidence','unknown']:raise ValueError('Unknown event')
    apply_event(store,flow,c,jid,mid,data,event,actor)
    c.execute("UPDATE messages SET job_id=?,status='applied' WHERE id=?",(jid,mid))
    return 'applied'

def scan_mailbox(store,flow,config):
    # Passwords are referenced by environment variable, never written into the job database.
    password=os.environ.get(config['password_env'],'')
    if not password:raise ValueError('Mailbox password environment variable not configured')
    mailbox=config['id']
    processed=0
    with imaplib.IMAP4_SSL(config['host'],int(config.get('port',993)),ssl_context=ssl.create_default_context(),timeout=int(config.get('timeout',45))) as mail:
        mail.login(config['username'],password)
        for folder in config.get('folders',['INBOX','Sent']):
            status,_=mail.select('"'+folder.replace('"','')+'"',readonly=True)
            if status!='OK':raise RuntimeError('Cannot select configured folder '+folder)
            response=mail.response('UIDVALIDITY')[1]
            validity=response[0].decode() if response and response[0] else ''
            if not validity:raise RuntimeError('Mailbox did not supply UIDVALIDITY')
            with store.connection() as c:checkpoint=c.execute('SELECT * FROM checkpoints WHERE mailbox=? AND folder=?',(mailbox,folder)).fetchone()
            if checkpoint and checkpoint['validity']==validity:start=checkpoint['uid']+1
            elif not checkpoint:
                # Default: monitor new mail from activation; explicit history import handles active jobs.
                uidnext=mail.response('UIDNEXT')[1]
                start=int(uidnext[0]) if uidnext and uidnext[0] else 1
                with store.transaction() as c:c.execute('INSERT OR REPLACE INTO checkpoints VALUES (?,?,?,?)',(mailbox,folder,validity,start-1))
            else:start=1 # UIDVALIDITY reset: rescan and deduplicate.
            status,values=mail.uid('search',None,f'UID {start}:*')
            if status!='OK':raise RuntimeError('IMAP UID search failed')
            uids=sorted(int(v) for v in (values[0] or b'').split() if int(v)>=start)
            for uid in uids[:int(config.get('batch_size',200))]:
                status,parts=mail.uid('fetch',str(uid),'(BODY.PEEK[])')
                if status!='OK':raise RuntimeError('IMAP fetch failed')
                raw=next((p[1] for p in parts if isinstance(p,tuple)),None)
                if raw is None:raise RuntimeError('Missing email body')
                ingest(store,flow,raw,mailbox,folder,validity,uid,mailbox_owner=config['id'] if parse_message(raw)['sender']==config['username'].lower() else '')
                with store.transaction() as c:c.execute('INSERT OR REPLACE INTO checkpoints VALUES (?,?,?,?)',(mailbox,folder,validity,uid))
                processed+=1
    return processed

def extract_attachment(store,digest,filename):
    with store.connection() as c:
        row=c.execute('SELECT data FROM extraction_cache WHERE hash=?',(digest,)).fetchone()
        if row:
            cached=json.loads(row[0])
            if cached.get('layout_version')==2 or Path(filename).suffix.lower()!='.pdf':return cached
    raw=store.blob('attachments',digest).read_bytes()
    suffix=Path(filename).suffix.lower()
    text='';fields={};tables=[]
    if suffix=='.pdf':
        import pymupdf
        from .document_facts import labelled_fields
        with pymupdf.open(stream=raw,filetype='pdf') as pdf:
            text='\n'.join(p.get_text() for p in pdf)[:40000]
            for page in list(pdf)[:10]:
                for key,value in labelled_fields([b[4] for b in page.get_text('blocks') if len(b)>6 and b[6]==0]).items():fields.setdefault(key,value)
                try:tables.extend(t.extract() for t in page.find_tables().tables)
                except Exception:pass
    elif suffix in ['.xlsx','.xlsm']:
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:text='\n'.join('\t'.join(str(v or '') for v in row) for sheet in book for row in sheet.iter_rows(values_only=True))[:40000]
        finally:book.close()
    elif suffix=='.docx':
        import docx
        doc=docx.Document(io.BytesIO(raw));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join('\t'.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    elif suffix in ['.txt','.csv']:text=raw.decode('utf-8',errors='replace')[:40000]
    result=dict(text=text[:40000],needs_vision=not bool(text.strip()),fields=fields,tables=tables,layout_version=2)
    with store.transaction() as c:c.execute('INSERT INTO extraction_cache VALUES (?,?) ON CONFLICT(hash) DO UPDATE SET data=excluded.data',(digest,json.dumps(result)))
    return result
