import os
import io
import csv
import re
import json
import base64
from flask import Flask, request, jsonify, render_template, send_file
from flask_cors import CORS
import pandas as pd
import openpyxl
import pymupdf
import docx

from extractor import parse_excel, parse_pdf, parse_image_file, parse_raw_text, detect_currency
from matcher import match_rfq_and_quotation, recalculate_prices, generate_accounting_tsv
from gemini_vision import (
    parse_rfq_gemini,
    parse_quote_gemini,
    parse_image_base64_gemini,
    parse_text_gemini,
    get_api_key,
    set_api_key
)

import sys
import smtplib
import imaplib
import ssl
import email
import email.mime
import email.mime.multipart
import email.mime.text
import email.mime.image
import email.mime.base
import email.encoders
import email.utils
import email.header
import email.message

BASE_DIR = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))

# Ensure email_automation directory is in sys.path
_ea_dir = os.path.join(BASE_DIR, 'email_automation')
if _ea_dir not in sys.path:
    sys.path.insert(0, _ea_dir)

try:
    from mailer import send_email_smtp, verify_smtp_connection
except Exception:
    try:
        from email_automation.mailer import send_email_smtp, verify_smtp_connection
    except Exception:
        send_email_smtp = None
        verify_smtp_connection = None

app = Flask(__name__, 
            template_folder=os.path.join(BASE_DIR, 'templates'), 
            static_folder=os.path.join(BASE_DIR, 'static'))
CORS(app)

USD_BASE_RATES = {
    'USD': 1.0000,
    'MYR': 4.4500,
    'EUR': 0.9200,
    'SGD': 1.3400,
    'GBP': 0.7850,
    'AUD': 1.5200,
    'CNY': 7.2300,
    'JPY': 154.5000,
    'AED': 3.6725
}

def generate_doc_previews(file_bytes, filename):
    """
    Generates rich visual and tabular preview data for any file format:
    PDF, Screenshots/Images, Excel, Word Docx, and Text/CSV.
    """
    fn = filename.lower()
    preview_data = {
        'filename': filename,
        'format': 'unknown',
        'pages': [],
        'tables': [],
        'text': ''
    }
    
    try:
        # 1. PDF Documents
        if fn.endswith('.pdf'):
            preview_data['format'] = 'pdf'
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            for i in range(min(len(doc), 10)): # up to 10 pages
                pix = doc[i].get_pixmap(dpi=180)
                b64 = base64.b64encode(pix.tobytes("png")).decode('utf-8')
                preview_data['pages'].append(f"data:image/png;base64,{b64}")
                
        # 2. Images / Screenshots
        elif any(fn.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp']):
            preview_data['format'] = 'image'
            b64 = base64.b64encode(file_bytes).decode('utf-8')
            mime = 'png' if fn.endswith('.png') else 'jpeg'
            preview_data['pages'].append(f"data:image/{mime};base64,{b64}")
            
        # 3. Excel Spreadsheets (.xlsx, .xls)
        elif fn.endswith('.xlsx') or fn.endswith('.xls'):
            preview_data['format'] = 'excel'
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                for sheet_name in wb.sheetnames[:5]: # up to 5 sheets
                    sheet = wb[sheet_name]
                    rows_data = []
                    for r in sheet.iter_rows(values_only=True):
                        if any(r):
                            clean_row = [str(c) if c is not None else "" for c in r]
                            rows_data.append(clean_row)
                    if rows_data:
                        preview_data['tables'].append({
                            'sheet_name': sheet_name,
                            'headers': rows_data[0] if len(rows_data) > 0 else [],
                            'rows': rows_data[1:100] if len(rows_data) > 1 else [] # up to 100 preview rows
                        })
            except Exception:
                try:
                    df_dict = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
                    for sheet_name, df in list(df_dict.items())[:5]:
                        rows_data = [[str(c) if pd.notna(c) else "" for c in r] for r in df.values]
                        preview_data['tables'].append({
                            'sheet_name': str(sheet_name),
                            'headers': [str(c) for c in df.columns],
                            'rows': rows_data[:100]
                        })
                except Exception:
                    pass
                    
        # 4. Word Documents (.docx, .doc)
        elif fn.endswith('.docx') or fn.endswith('.doc'):
            preview_data['format'] = 'docx'
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            preview_data['text'] = "\n\n".join(paragraphs[:30])
            
            for t_idx, table in enumerate(doc.tables[:5]):
                t_rows = []
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells]
                    if any(row_cells):
                        t_rows.append(row_cells)
                if t_rows:
                    preview_data['tables'].append({
                        'sheet_name': f"Table {t_idx+1}",
                        'headers': t_rows[0],
                        'rows': t_rows[1:]
                    })
                    
        # 5. CSV or Text files
        elif fn.endswith('.csv') or fn.endswith('.txt'):
            preview_data['format'] = 'text'
            text_str = file_bytes.decode('utf-8', errors='ignore')
            lines = [l.strip() for l in text_str.split('\n') if l.strip()]
            preview_data['text'] = "\n".join(lines[:100])
            if fn.endswith('.csv'):
                import csv
                reader = csv.reader(io.StringIO(text_str))
                c_rows = [r for r in reader if any(r)]
                if c_rows:
                    preview_data['tables'].append({
                        'sheet_name': 'CSV Data',
                        'headers': c_rows[0],
                        'rows': c_rows[1:100]
                    })
    except Exception as e:
        print(f"Preview generation exception: {e}")
        
    return preview_data

@app.route('/')
def index():
    has_gemini = bool(get_api_key())
    return render_template('index.html', default_usd_rates=USD_BASE_RATES, has_gemini=has_gemini)

@app.route('/api/config/api-key', methods=['GET', 'POST'])
def handle_api_key_config():
    if request.method == 'POST':
        data = request.json or {}
        new_key = data.get('api_key', '').strip()
        set_api_key(new_key)
        return jsonify({'success': True, 'has_key': bool(new_key)})
    else:
        current_key = get_api_key()
        masked_key = f"{current_key[:6]}...{current_key[-4:]}" if len(current_key) > 10 else ("Configured" if current_key else "")
        return jsonify({
            'success': True,
            'has_key': bool(current_key),
            'masked_key': masked_key
        })

@app.route('/api/parse-rfq', methods=['POST'])
def handle_parse_rfq():
    try:
        items = []
        source_engine = "local"
        preview_data = None
        
        if 'file' in request.files:
            file = request.files['file']
            filename = file.filename.lower()
            file_bytes = file.read()
            preview_data = generate_doc_previews(file_bytes, file.filename)
            
            # 1. Try Gemini Vision AI
            if get_api_key():
                try:
                    gemini_items = parse_rfq_gemini(file_bytes, filename=filename)
                    if gemini_items and len(gemini_items) > 0:
                        items = gemini_items
                        source_engine = "gemini"
                except Exception as ex:
                    print(f"Gemini RFQ parsing exception: {ex}")
            
            # 2. Local Fallback
            if not items:
                if filename.endswith('.xlsx') or filename.endswith('.xls'):
                    items = parse_excel(file_bytes)
                elif filename.endswith('.pdf'):
                    items, _ = parse_pdf(file_bytes)
                elif any(filename.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp']):
                    items, _ = parse_image_file(file_bytes)
                elif filename.endswith('.docx') or filename.endswith('.doc'):
                    doc = docx.Document(io.BytesIO(file_bytes))
                    full_text = "\n".join([p.text for p in doc.paragraphs if p.text])
                    items, _ = parse_raw_text(full_text)
                else:
                    return jsonify({'success': False, 'error': 'Unsupported file format'}), 400
                source_engine = "local"
                
        elif request.is_json and 'text' in request.json:
            text = request.json.get('text', '')
            preview_data = {'filename': 'Pasted Text', 'format': 'text', 'pages': [], 'tables': [], 'text': text}
            if get_api_key():
                try:
                    gemini_items = parse_text_gemini(text, doc_type="rfq")
                    if gemini_items:
                        items = gemini_items
                        source_engine = "gemini"
                except Exception:
                    pass
            if not items:
                items, _ = parse_raw_text(text)
                source_engine = "local"
            
        elif request.is_json and 'image_base64' in request.json:
            b64_data = request.json.get('image_base64', '')
            img_url = b64_data if b64_data.startswith('data:') else f"data:image/png;base64,{b64_data}"
            preview_data = {'filename': 'Pasted Screenshot', 'format': 'image', 'pages': [img_url], 'tables': [], 'text': ''}
            if get_api_key():
                try:
                    gemini_items = parse_image_base64_gemini(b64_data, doc_type="rfq")
                    if gemini_items:
                        items = gemini_items
                        source_engine = "gemini"
                except Exception as ex:
                    print(f"Gemini RFQ image_base64 exception: {ex}")
            if not items:
                if ',' in b64_data:
                    b64_data = b64_data.split(',')[1]
                img_bytes = base64.b64decode(b64_data)
                items, _ = parse_image_file(img_bytes)
                source_engine = "local"
            
        else:
            return jsonify({'success': False, 'error': 'No file or data provided'}), 400
            
        return jsonify({
            'success': True,
            'count': len(items),
            'items': items,
            'engine': source_engine,
            'preview_data': preview_data
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/parse-quotation', methods=['POST'])
def handle_parse_quotation():
    try:
        items = []
        currency = 'MYR'
        source_engine = "local"
        preview_data = None
        
        if 'file' in request.files:
            file = request.files['file']
            filename = file.filename.lower()
            file_bytes = file.read()
            preview_data = generate_doc_previews(file_bytes, file.filename)
            
            # 1. Try Gemini Vision AI
            if get_api_key():
                try:
                    gemini_items = parse_quote_gemini(file_bytes, filename=filename)
                    if gemini_items and len(gemini_items) > 0:
                        items = gemini_items
                        source_engine = "gemini"
                        for it in items:
                            if it.get('currency'):
                                currency = it['currency'].upper()
                                break
                except Exception as ex:
                    print(f"Gemini Quote parsing exception: {ex}")
                    
            # 2. Local Fallback
            if not items:
                if filename.endswith('.xlsx') or filename.endswith('.xls'):
                    items = parse_excel(file_bytes)
                    currency = 'MYR'
                elif filename.endswith('.pdf'):
                    items, currency = parse_pdf(file_bytes)
                elif any(filename.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp']):
                    items, currency = parse_image_file(file_bytes)
                elif filename.endswith('.docx') or filename.endswith('.doc'):
                    doc = docx.Document(io.BytesIO(file_bytes))
                    full_text = "\n".join([p.text for p in doc.paragraphs if p.text])
                    items, currency = parse_raw_text(full_text)
                else:
                    return jsonify({'success': False, 'error': 'Unsupported file format'}), 400
                source_engine = "local"
                
        elif request.is_json and 'text' in request.json:
            text = request.json.get('text', '')
            preview_data = {'filename': 'Pasted Text', 'format': 'text', 'pages': [], 'tables': [], 'text': text}
            if get_api_key():
                try:
                    gemini_items = parse_text_gemini(text, doc_type="quote")
                    if gemini_items:
                        items = gemini_items
                        source_engine = "gemini"
                        for it in items:
                            if it.get('currency'):
                                currency = it['currency'].upper()
                                break
                except Exception:
                    pass
            if not items:
                items, currency = parse_raw_text(text)
                source_engine = "local"
            
        elif request.is_json and 'image_base64' in request.json:
            b64_data = request.json.get('image_base64', '')
            img_url = b64_data if b64_data.startswith('data:') else f"data:image/png;base64,{b64_data}"
            preview_data = {'filename': 'Pasted Screenshot', 'format': 'image', 'pages': [img_url], 'tables': [], 'text': ''}
            if get_api_key():
                try:
                    gemini_items = parse_image_base64_gemini(b64_data, doc_type="quote")
                    if gemini_items:
                        items = gemini_items
                        source_engine = "gemini"
                        for it in items:
                            if it.get('currency'):
                                currency = it['currency'].upper()
                                break
                except Exception as ex:
                    print(f"Gemini Quote image_base64 exception: {ex}")
            if not items:
                if ',' in b64_data:
                    b64_data = b64_data.split(',')[1]
                img_bytes = base64.b64decode(b64_data)
                items, currency = parse_image_file(img_bytes)
                source_engine = "local"
            
        else:
            return jsonify({'success': False, 'error': 'No file or data provided'}), 400
            
        return jsonify({
            'success': True,
            'count': len(items),
            'currency': currency,
            'items': items,
            'engine': source_engine,
            'preview_data': preview_data
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/match', methods=['POST'])
def handle_match():
    try:
        data = request.json or {}
        rfq_items = data.get('rfq_items', [])
        quote_items = data.get('quote_items', [])
        fx_rate = float(data.get('fx_rate', 1.0))
        markup_percent = float(data.get('markup_percent', 25.0))
        supplier_currency = data.get('supplier_currency', 'MYR')
        target_currency = data.get('target_currency', 'MYR')
        mode = data.get('mode', 'quotation')
        
        matched_rows = match_rfq_and_quotation(
            rfq_items=rfq_items,
            quote_items=quote_items,
            fx_rate=fx_rate,
            markup_percent=markup_percent,
            supplier_currency=supplier_currency,
            target_currency=target_currency,
            mode=mode
        )
        
        tsv_str = generate_accounting_tsv(matched_rows, target_currency=target_currency, mode=mode)
        
        return jsonify({
            'success': True,
            'count': len(matched_rows),
            'items': matched_rows,
            'target_currency': target_currency,
            'supplier_currency': supplier_currency,
            'mode': mode,
            'tsv': tsv_str
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/recalculate', methods=['POST'])
def handle_recalculate():
    try:
        data = request.json or {}
        items = data.get('items', [])
        fx_rate = float(data.get('fx_rate', 1.0))
        target_currency = data.get('target_currency', 'MYR')
        mode = data.get('mode', 'quotation')
        global_markup = data.get('global_markup', None)
        if global_markup is not None:
            global_markup = float(global_markup)
            
        global_adjustment_type = data.get('global_adjustment_type', 'percent')
        global_markup_amount = data.get('global_markup_amount', None)
        if global_markup_amount is not None:
            global_markup_amount = float(global_markup_amount)
            
        updated_items = recalculate_prices(
            items, 
            fx_rate=fx_rate, 
            global_markup=global_markup, 
            target_currency=target_currency,
            mode=mode,
            global_adjustment_type=global_adjustment_type,
            global_markup_amount=global_markup_amount
        )
        tsv_str = generate_accounting_tsv(updated_items, target_currency=target_currency, mode=mode)
        
        return jsonify({
            'success': True,
            'items': updated_items,
            'target_currency': target_currency,
            'mode': mode,
            'tsv': tsv_str
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/export-excel', methods=['POST'])
def handle_export_excel():
    try:
        data = request.json or {}
        items = data.get('items', [])
        target_currency = data.get('target_currency', 'MYR')
        mode = data.get('mode', 'quotation')
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PO Import" if mode == "po" else "Quotation Import"
        
        price_header = f"Unit Cost ({target_currency})" if mode == "po" else f"Price ({target_currency})"
        headers = ["#", "Item Code", "Description", "Quantity", "Unit", price_header]
        ws.append(headers)
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            header_color = "065F46" if mode == "po" else "1E293B"
            cell.fill = openpyxl.styles.PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
            
        for idx, item in enumerate(items, start=1):
            price_val = float(item.get('price_target', item.get('price_myr', 0.0)))
            ws.append([
                idx,
                item.get('item_code', ''),
                item.get('description', ''),
                item.get('quantity', 1),
                item.get('unit', 'PCS'),
                price_val
            ])
            
            price_cell = ws.cell(row=idx + 1, column=6)
            price_cell.number_format = '#,##0.00'
            
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        prefix = "AP_Merin_Purchase_Order" if mode == "po" else "AP_Merin_Quotation"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{prefix}_{target_currency}.xlsx"
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/export-csv', methods=['POST'])
def handle_export_csv():
    try:
        data = request.json or {}
        items = data.get('items', [])
        target_currency = data.get('target_currency', 'MYR')
        mode = data.get('mode', 'quotation')
        
        price_header = f"Unit Cost ({target_currency})" if mode == "po" else f"Price ({target_currency})"
        rows = []
        for idx, item in enumerate(items, start=1):
            price_val = float(item.get('price_target', item.get('price_myr', 0.0)))
            rows.append({
                '#': idx,
                'Item Code': item.get('item_code', ''),
                'Description': item.get('description', ''),
                'Quantity': item.get('quantity', 1),
                'Unit': item.get('unit', 'PCS'),
                price_header: price_val
            })
            
        df = pd.DataFrame(rows)
        output = io.StringIO()
        df.to_csv(output, index=False)
        
        mem = io.BytesIO()
        mem.write(output.getvalue().encode('utf-8-sig'))
        mem.seek(0)
        
        prefix = "AP_Merin_Purchase_Order" if mode == "po" else "AP_Merin_Quotation"
        return send_file(
            mem,
            mimetype="text/csv",
            as_attachment=True,
            download_name=f"{prefix}_{target_currency}.csv"
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/.netlify/functions/send-email', methods=['POST', 'OPTIONS'])
def netlify_send_email_proxy():
    if request.method == 'OPTIONS':
        return '', 200
    try:
        import tempfile
        import base64
        
        send_fn = send_email_smtp
        if send_fn is None:
            from mailer import send_email_smtp as send_fn

        data = request.json or {}
        recipients = data.get('recipients', [])
        subject = data.get('subject', '').strip()
        html_body = data.get('html_body', '').strip()
        cc_emails = data.get('cc_emails', [])
        smtp_cfg = data.get('smtp_config', {})
        raw_attachments = data.get('attachments', [])
        
        host = smtp_cfg.get('smtp_host', 'smtppro47.mschosting.network')
        port = int(smtp_cfg.get('smtp_port', 465))
        username = smtp_cfg.get('username', 'project@apmerinservices.com.my')
        password = smtp_cfg.get('password', '')
        from_email = smtp_cfg.get('sender_email', username)
        from_name = smtp_cfg.get('sender_name', 'AP Merin Services')

        # Convert raw base64 dataUrl attachments preserving exact original filenames
        temp_dir = tempfile.mkdtemp(prefix="ap_att_dir_")
        temp_file_items = []
        try:
            for att in raw_attachments:
                data_url = att.get('dataUrl', '')
                raw_filename = (att.get('name') or 'attachment.bin').strip()
                # Sanitize name to avoid path traversal while preserving full original name and extension
                clean_filename = os.path.basename(raw_filename) or 'attachment.bin'
                if 'base64,' in data_url:
                    b64_str = data_url.split('base64,')[1]
                    file_bytes = base64.b64decode(b64_str)
                    dest_path = os.path.join(temp_dir, clean_filename)
                    counter = 1
                    base_stem, ext = os.path.splitext(clean_filename)
                    while os.path.exists(dest_path):
                        dest_path = os.path.join(temp_dir, f"{base_stem}_{counter}{ext}")
                        counter += 1
                    with open(dest_path, 'wb') as f:
                        f.write(file_bytes)
                    temp_file_items.append((dest_path, clean_filename))

            sent_count = 0
            results = []
            for item in recipients:
                if isinstance(item, dict):
                    comp = item.get('company', '')
                    emails = item.get('emails', [])
                elif isinstance(item, list):
                    comp = ''
                    emails = item
                else:
                    comp = ''
                    emails = [str(item)]
                
                clean_emails = [e.strip() for e in emails if '@' in e]
                if not clean_emails:
                    continue
                
                success, msg = send_fn(
                    host=host,
                    port=port,
                    username=username,
                    password=password,
                    sender_email=from_email,
                    sender_name=from_name,
                    recipient_email=clean_emails,
                    subject=subject,
                    html_body=html_body,
                    cc_emails=cc_emails,
                    file_attachments=temp_file_items,
                    use_ssl=(port == 465),
                    use_tls=(port == 587)
                )
                if success:
                    sent_count += 1
                    results.append({'company': comp, 'emails': clean_emails, 'status': 'sent', 'message': 'Sent successfully'})
                else:
                    results.append({'company': comp, 'emails': clean_emails, 'status': 'failed', 'message': msg})
                    
            is_overall_success = (sent_count > 0)
            primary_msg = f'Dispatched {sent_count} of {len(recipients)} email(s)!' if is_overall_success else (results[0]['message'] if results else 'Failed to dispatch email.')
            return jsonify({
                'success': is_overall_success,
                'sentCount': sent_count,
                'total': len(recipients),
                'results': results,
                'message': primary_msg
            })
        finally:
            if temp_dir and os.path.exists(temp_dir):
                try:
                    shutil.rmtree(temp_dir, ignore_errors=True)
                except Exception:
                    pass
    except Exception as e:
        return jsonify({'success': False, 'message': f"Server error: {str(e)}"}), 500


# =========================================================================
# QUOTATION COMPARISON MAKER ENDPOINTS
# =========================================================================
from quotation_comparator import (
    search_imap_emails,
    search_outlook_desktop_emails,
    parse_rfc822_email,
    smart_deduplicate_and_organize,
    analyze_email_threads_and_lifecycle,
    process_supplier_package,
    compile_side_by_side_comparison,
    export_comparison_to_excel,
    extract_single_quote_with_gemini,
    extract_single_quote_local_fallback,
    extract_baseline_rfq,
    generate_source_preview_data,
    find_initial_rfq_email,
    extract_sender_identity,
    extract_rfq_fingerprint,
    is_email_relevant_to_rfq
)

comparator_scanned_cache = {}

@app.route('/api/comparator/scan-emails', methods=['POST'])
def handle_comparator_scan_emails():
    """
    STAGE 1: Searches mailbox and discovers all matching email threads.
    Analyzes initial RFQ originator (ignoring CC), internal staff follow-ups,
    and categorizes each supplier's communication lifecycle without consuming AI extraction tokens.
    """
    try:
        data = request.json or {}
        subject_query = data.get('subject_query', '').strip()
        mailbox_user = (data.get('mailbox_user') or '').strip() or "project@apmerinservices.com.my"

        if not subject_query:
            return jsonify({'success': False, 'error': 'Please enter an RFQ subject or keyword to search.'}), 400

        raw_emails = []
        source_used = "imap"

        # 1. Try IMAP search first (primary authoritative mailbox with shared 'From AP Merin' folder)
        try:
            imap_raw = search_imap_emails(subject_query, max_results=12)
            for rb in imap_raw:
                parsed_em = parse_rfc822_email(rb)
                raw_emails.append(parsed_em)
            source_used = "imap"
        except Exception as ex:
            print(f"IMAP mailbox search info: {ex}")

        # 2. Try Outlook Desktop if IMAP returned no emails or failed
        if not raw_emails:
            try:
                outlook_emails = search_outlook_desktop_emails(subject_query, max_results=12, mailbox_user=mailbox_user)
                if outlook_emails:
                    raw_emails.extend(outlook_emails)
                    source_used = "outlook"
            except Exception as oex:
                print(f"Outlook desktop search info: {oex}")

        if not raw_emails:
            return jsonify({
                'success': False,
                'error': f'No emails found matching "{subject_query}". Please check the RFQ subject or reference number.'
            }), 404

        # 3. Analyze thread lifecycle & sender attribution
        thread_info = analyze_email_threads_and_lifecycle(raw_emails, mailbox_user=mailbox_user, target_rfq_query=subject_query)

        # Cache in memory for instantaneous Stage 2 extraction
        query_key = subject_query.lower().strip()
        comparator_scanned_cache[query_key] = {
            'raw_emails': raw_emails,
            'thread_info': thread_info,
            'source_used': source_used
        }

        return jsonify({
            'success': True,
            'subject_query': subject_query,
            'source_used': source_used,
            'initial_rfq': thread_info.get('initial_rfq', {}),
            'supplier_threads': thread_info.get('supplier_threads', []),
            'discovered_emails': thread_info.get('discovered_emails', []),
            'dedup_stats': thread_info.get('dedup_stats', {}),
            'total_emails_found': thread_info.get('total_emails_found', 0),
            'total_suppliers_found': thread_info.get('total_suppliers_found', 0)
        })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/comparator/fetch-quote-doc', methods=['POST'])
def handle_comparator_fetch_quote_doc():
    """
    Fetches quotation attachments for approved suppliers.
    Uses cached emails in memory or IMAP fetch directly.
    Generates both base64 raw binary and PyMuPDF preview pages for instantaneous rendering.
    """
    try:
        data = request.json or {}
        items = data.get('items', [data])
        documents = []

        for it in items:
            filename = (it.get('filename') or '').strip()
            folder = it.get('folder', 'INBOX')
            msg_id = it.get('msg_id')
            supp_key = it.get('supplier_key', '')

            raw_bytes = None
            clean_search_name = filename.strip('"\' ')

            # 1. Check cached emails in memory first
            for cache_val in comparator_scanned_cache.values():
                for em in cache_val.get('raw_emails', []):
                    for fn, fb in em.get('attachments', []):
                        fn_clean = fn.strip('"\' ')
                        if fn_clean.lower() == clean_search_name.lower() or clean_search_name.lower() in fn_clean.lower() or fn_clean.lower() in clean_search_name.lower():
                            raw_bytes = fb
                            break
                    if raw_bytes:
                        break
                if raw_bytes:
                    break

            # 2. Fallback: Search IMAP directly in Python if not found in cache
            if not raw_bytes and clean_search_name:
                try:
                    imap_raw = search_imap_emails(clean_search_name, max_results=5)
                    for rb in imap_raw:
                        parsed_em = parse_rfc822_email(rb)
                        for fn, fb in parsed_em.get('attachments', []):
                            fn_clean = fn.strip('"\' ')
                            if fn_clean.lower() == clean_search_name.lower() or clean_search_name.lower() in fn_clean.lower():
                                raw_bytes = fb
                                break
                        if raw_bytes:
                            break
                except Exception as ex:
                    print(f"IMAP fallback fetch error in fetch-quote-doc: {ex}")

            base64_str = ''
            preview_pages = []
            if raw_bytes:
                base64_str = base64.b64encode(raw_bytes).decode('utf-8')
                try:
                    p_data = generate_source_preview_data(doc_bytes=raw_bytes, filename=filename)
                    preview_pages = p_data.get('preview_pages', [])
                except Exception as e:
                    print(f"Preview pages generation error in fetch-quote-doc: {e}")

            fn_l = filename.lower()
            if fn_l.endswith('.pdf'):
                eff_mime = 'application/pdf'
            elif fn_l.endswith(('.xlsx', '.xls')):
                eff_mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            elif fn_l.endswith('.csv'):
                eff_mime = 'text/csv'
            elif fn_l.endswith(('.docx', '.doc')):
                eff_mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            elif any(fn_l.endswith(img_ext) for img_ext in ['.png', '.jpg', '.jpeg', '.gif', '.webp']):
                eff_mime = 'image/png' if fn_l.endswith('.png') else 'image/jpeg'
            else:
                eff_mime = 'application/octet-stream'

            documents.append({
                'supplier_key': supp_key,
                'filename': filename,
                'msg_id': msg_id,
                'folder': folder,
                'mime_type': eff_mime,
                'base64': base64_str,
                'preview_pages': preview_pages
            })

        return jsonify({
            'success': True,
            'count': len(documents),
            'documents': documents
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/comparator/proceed-extraction', methods=['POST'])
def handle_comparator_proceed_extraction():
    """
    STAGE 2: Runs AI extraction on the user-approved suppliers and compiles
    the side-by-side comparison matrix with baseline verification.
    """
    try:
        data = request.json or {}
        subject_query = data.get('subject_query', '').strip()
        approved_keys = data.get('approved_supplier_keys', []) # list of supplier_key strings
        target_currency = data.get('target_currency', 'MYR')
        fx_rates = data.get('fx_rates', USD_BASE_RATES)

        query_key = subject_query.lower().strip()
        cached = comparator_scanned_cache.get(query_key)

        raw_emails = []
        source_used = "imap"
        initial_rfq = {}

        if cached:
            raw_emails = cached.get('raw_emails', [])
            source_used = cached.get('source_used', 'imap')
            initial_rfq = cached.get('thread_info', {}).get('initial_rfq', {})
        else:
            # Fallback re-fetch if cache expired
            try:
                imap_raw = search_imap_emails(subject_query, max_results=12)
                for rb in imap_raw:
                    raw_emails.append(parse_rfc822_email(rb))
            except Exception as ex:
                return jsonify({'success': False, 'error': f"Failed fetching emails: {str(ex)}"}), 500

        if not raw_emails:
            return jsonify({'success': False, 'error': 'No emails found to extract.'}), 404

        if subject_query:
            fp = extract_rfq_fingerprint(subject_query)
            rel_emails = [em for em in raw_emails if is_email_relevant_to_rfq(em, fp)[0]]
            if rel_emails:
                raw_emails = rel_emails

        mailbox_user = (data.get('mailbox_user') or '').strip() or "project@apmerinservices.com.my"
        supplier_packages, dedup_stats, outgoing_emails = smart_deduplicate_and_organize(raw_emails, user_email=mailbox_user)

        # Filter by approved suppliers if user deselected any
        if approved_keys:
            approved_set = set(approved_keys)
            filtered_packages = [pkg for pkg in supplier_packages if pkg.get('supplier_key') in approved_set or pkg.get('sender_name') in approved_set]
            if filtered_packages:
                supplier_packages = filtered_packages

        gemini_key = get_api_key()

        # Extract Baseline RFQ items
        baseline_rfq = None
        try:
            baseline_rfq = extract_baseline_rfq(outgoing_emails=outgoing_emails, gemini_key=gemini_key, all_emails=raw_emails)
        except Exception as bex:
            print(f"Baseline RFQ extraction notice: {bex}")

        # Extract quotes from each approved supplier package
        supplier_quotes = []
        for pkg in supplier_packages:
            try:
                quote = process_supplier_package(pkg, gemini_key=gemini_key)
                if quote and (quote.get('items') or quote.get('is_regret') or quote.get('is_pending')):
                    supplier_quotes.append(quote)
            except Exception as e:
                print(f"Extraction error for supplier {pkg.get('supplier_key')}: {e}")

        if not supplier_quotes:
            return jsonify({
                'success': False,
                'error': 'Could not extract quotation data from the approved supplier emails.'
            }), 422

        if not initial_rfq and outgoing_emails:
            init_em = find_initial_rfq_email(outgoing_emails, raw_emails)
            if init_em:
                s_name, s_email = extract_sender_identity(init_em.get('from', ''), init_em.get('body', ''))
                initial_rfq = {
                    "sender_name": s_name,
                    "sender_email": s_email,
                    "from": init_em.get('from', ''),
                    "to": init_em.get('to', ''),
                    "cc": init_em.get('cc', ''),
                    "date": init_em.get('date', ''),
                    "subject": init_em.get('subject', ''),
                    "body_text": init_em.get('body', ''),
                    "html_body": init_em.get('html_body', ''),
                    "html_body_resolved": init_em.get('html_body_resolved', ''),
                    "images": init_em.get('images', []),
                    "rich_attachments": init_em.get('rich_attachments', [])
                }

        # Compile comparison matrix
        comparison_data = compile_side_by_side_comparison(
            supplier_quotes,
            target_currency=target_currency,
            fx_rates=fx_rates,
            baseline_rfq=baseline_rfq,
            initial_rfq=initial_rfq
        )
        comparison_data['dedup_stats'] = dedup_stats
        comparison_data['source_used'] = source_used
        comparison_data['subject_query'] = subject_query

        return jsonify(comparison_data)

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/comparator/search-and-extract', methods=['POST'])
def handle_comparator_search_and_extract():
    try:
        data = request.json or {}
        subject_query = data.get('subject_query', '').strip()
        target_currency = data.get('target_currency', 'MYR')
        fx_rates = data.get('fx_rates', USD_BASE_RATES)

        if not subject_query:
            return jsonify({'success': False, 'error': 'Please enter an RFQ subject or keyword to search.'}), 400

        raw_emails = []
        source_used = "imap"

        # 1. Try IMAP search first
        try:
            imap_raw = search_imap_emails(subject_query, max_results=20)
            for rb in imap_raw:
                parsed_em = parse_rfc822_email(rb)
                raw_emails.append(parsed_em)
            source_used = "imap"
        except Exception as ex:
            print(f"IMAP search note: {ex}")

        # 2. Try Outlook Desktop if IMAP returned no emails or failed
        if not raw_emails:
            try:
                outlook_emails = search_outlook_desktop_emails(subject_query, max_results=20)
                if outlook_emails:
                    raw_emails.extend(outlook_emails)
                    source_used = "outlook"
            except Exception as oex:
                print(f"Outlook search note: {oex}")

        if not raw_emails:
            return jsonify({
                'success': False,
                'error': f'No emails found matching subject "{subject_query}". Please check the subject name or reference number.'
            }), 404

        # 3. Smart Deduplication & Token Optimization
        supplier_packages, dedup_stats, outgoing_emails = smart_deduplicate_and_organize(raw_emails)

        if not supplier_packages and not outgoing_emails:
            return jsonify({
                'success': False,
                'error': 'No supplier replies found (only unparseable emails were found).'
            }), 404

        gemini_key = get_api_key()

        # 4. Extract Baseline RFQ items (what AP Merin sent out)
        baseline_rfq = None
        try:
            baseline_rfq = extract_baseline_rfq(outgoing_emails=outgoing_emails, gemini_key=gemini_key, all_emails=raw_emails)
        except Exception as bex:
            print(f"Baseline RFQ extraction notice: {bex}")

        # 5. Extract Quotes from each Supplier Package
        supplier_quotes = []

        for pkg in supplier_packages:
            try:
                quote = process_supplier_package(pkg, gemini_key=gemini_key)
                if quote and (quote.get('items') or quote.get('is_regret') or quote.get('is_pending')):
                    supplier_quotes.append(quote)
            except Exception as e:
                print(f"Failed extracting quote from supplier {pkg.get('supplier_key')}: {e}")

        if not supplier_quotes:
            return jsonify({
                'success': False,
                'error': 'Could not extract quotation data from the found supplier emails/attachments. Please verify files or try uploading manually.'
            }), 422

        # 6. Compile Side-by-Side Comparison Matrix with Baseline Verification
        initial_rfq = (baseline_rfq.get('baseline_email') if baseline_rfq else None) or {}
        if not initial_rfq and outgoing_emails:
            init_em = find_initial_rfq_email(outgoing_emails, raw_emails)
            if init_em:
                s_name, s_email = extract_sender_identity(init_em.get('from', ''), init_em.get('body', ''))
                initial_rfq = {
                    "sender_name": s_name,
                    "sender_email": s_email,
                    "from": init_em.get('from', ''),
                    "to": init_em.get('to', ''),
                    "cc": init_em.get('cc', ''),
                    "date": init_em.get('date', ''),
                    "subject": init_em.get('subject', ''),
                    "body_text": init_em.get('body', ''),
                    "html_body": init_em.get('html_body', ''),
                    "html_body_resolved": init_em.get('html_body_resolved', ''),
                    "images": init_em.get('images', []),
                    "rich_attachments": init_em.get('rich_attachments', [])
                }

        comparison_data = compile_side_by_side_comparison(
            supplier_quotes, 
            target_currency=target_currency, 
            fx_rates=fx_rates, 
            baseline_rfq=baseline_rfq,
            initial_rfq=initial_rfq
        )
        comparison_data['dedup_stats'] = dedup_stats
        comparison_data['source_used'] = source_used
        comparison_data['subject_query'] = subject_query

        return jsonify(comparison_data)

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/comparator/extract-manual', methods=['POST'])
def handle_comparator_extract_manual():
    try:
        target_currency = request.form.get('target_currency', 'MYR')
        gemini_key = get_api_key()
        supplier_quotes = []

        # Optional baseline RFQ file
        baseline_rfq = None
        if 'baseline_file' in request.files:
            b_file = request.files['baseline_file']
            if b_file and b_file.filename:
                bf_bytes = b_file.read()
                baseline_rfq = extract_baseline_rfq(uploaded_baseline_doc=(b_file.filename, bf_bytes), gemini_key=gemini_key)

        # Uploaded files
        files = request.files.getlist('files')
        for f in files:
            fn = f.filename
            fb = f.read()
            quote = None
            if gemini_key:
                try:
                    quote = extract_single_quote_with_gemini(doc_bytes=fb, filename=fn)
                except Exception:
                    pass
            if not quote:
                quote = extract_single_quote_local_fallback(doc_bytes=fb, filename=fn, sender_name=os.path.splitext(fn)[0])

            if quote and quote.get('items'):
                quote['source_file'] = fn
                quote['source_type'] = "Attachment"
                quote['source_preview'] = generate_source_preview_data(doc_bytes=fb, filename=fn)
                supplier_quotes.append(quote)

        # Pasted text / email body
        pasted_text = request.form.get('pasted_text', '').strip()
        if pasted_text:
            quote = None
            if gemini_key:
                try:
                    quote = extract_single_quote_with_gemini(email_text=pasted_text)
                except Exception:
                    pass
            if not quote:
                quote = extract_single_quote_local_fallback(email_text=pasted_text, sender_name="Pasted Quotation")
            if quote and quote.get('items'):
                quote['source_file'] = "Pasted Text"
                quote['source_type'] = "Email Body"
                quote['source_preview'] = generate_source_preview_data(email_info={"body_text": pasted_text, "sender_name": "Pasted Quotation"})
                supplier_quotes.append(quote)

        if not supplier_quotes:
            return jsonify({'success': False, 'error': 'No valid quote items could be extracted from the uploaded files or text.'}), 400

        comparison_data = compile_side_by_side_comparison(supplier_quotes, target_currency=target_currency, baseline_rfq=baseline_rfq)
        return jsonify(comparison_data)

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/comparator/recent-rfqs', methods=['GET'])
def handle_comparator_recent_rfqs():
    try:
        import imaplib, ssl, email
        from quotation_comparator import decode_mime_words, get_default_email_config
        cfg = get_default_email_config()
        host = cfg.get("imap_host") or cfg.get("smtp_host", "smtppro47.mschosting.network")
        username = cfg.get("username", "")
        password = cfg.get("password", "")

        if not username or not password:
            return jsonify({'success': True, 'rfqs': []})

        context = ssl.create_default_context()
        context.check_hostname = False
        context.verify_mode = ssl.CERT_NONE

        mail = imaplib.IMAP4_SSL(host, 993, ssl_context=context, timeout=10)
        mail.login(username, password)
        mail.select("INBOX", readonly=True)
        status, data = mail.search(None, 'SUBJECT', 'RFQ')
        msg_ids = data[0].split() if status == 'OK' and data[0] else []
        recent_ids = msg_ids[-25:] if len(msg_ids) > 25 else msg_ids
        recent_ids.reverse()

        rfqs = []
        seen = set()
        for mid in recent_ids:
            try:
                status, mdata = mail.fetch(mid, '(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE)])')
                for part in mdata:
                    if isinstance(part, tuple):
                        h_msg = email.message_from_bytes(part[1])
                        sub = decode_mime_words(h_msg.get('Subject', ''))
                        clean_sub = re.sub(r'^(Re:\s*|Fwd:\s*|On\s+Check\s+Our\s+Ref\.\s*no\.\s*[A-Z0-9\-]+\s*)+', '', sub, flags=re.IGNORECASE).strip()
                        if clean_sub and clean_sub not in seen:
                            seen.add(clean_sub)
                            rfqs.append(clean_sub)
            except Exception:
                pass
            if len(rfqs) >= 8:
                break
        mail.logout()
        return jsonify({'success': True, 'rfqs': rfqs})
    except Exception as e:
        print(f"Recent RFQs error: {e}")
        return jsonify({'success': True, 'rfqs': []})


@app.route('/api/comparator/export-excel', methods=['POST'])
def handle_comparator_export_excel():
    try:
        data = request.json or {}
        excel_bytes = export_comparison_to_excel(data)
        
        target_currency = data.get('target_currency', 'MYR')
        out_stream = io.BytesIO(excel_bytes)
        out_stream.seek(0)
        
        return send_file(
            out_stream,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"AP_Merin_Quotation_Comparison_{target_currency}.xlsx"
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# =========================================================================
# CROSS-DEVICE DATA SYNC & PERSISTENCE ENDPOINTS
# =========================================================================
import datetime
import sync_manager

# -------------------------------------------------------------------------
# Persistent Authentication Session Endpoints (Remember-Me / Auto-Login)
# -------------------------------------------------------------------------
@app.route('/api/auth/session', methods=['GET', 'POST'])
def handle_auth_session():
    try:
        if request.method == 'POST':
            data = request.json or {}
            user = data.get('user', {})
            remember_me = bool(data.get('remember_me', True))
            sess = sync_manager.save_auth_session(user, remember_me=remember_me)
            return jsonify({'success': True, 'session': sess})
        else:
            sess = sync_manager.load_auth_session()
            if sess and sess.get('authenticated'):
                return jsonify({'success': True, 'authenticated': True, 'session': sess})
            return jsonify({'success': True, 'authenticated': False, 'session': None})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/auth/logout', methods=['POST'])
def handle_auth_logout():
    try:
        sync_manager.clear_auth_session()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

# -------------------------------------------------------------------------
# User Profiles & Signatures Persistence Endpoints
# -------------------------------------------------------------------------
@app.route('/api/settings/users', methods=['GET', 'POST'])
def handle_settings_users():
    try:
        if request.method == 'POST':
            data = request.json or {}
            users = data.get('users')
            admin = data.get('admin')
            active_user_id = data.get('active_user_id')
            bundle = sync_manager.save_user_profiles_bundle(users=users, admin=admin, active_user_id=active_user_id)
            return jsonify({'success': True, 'bundle': bundle})
        else:
            bundle = sync_manager.load_user_profiles_bundle()
            return jsonify({'success': True, 'bundle': bundle})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/settings/signatures', methods=['GET', 'POST'])
def handle_settings_signatures():
    try:
        if request.method == 'POST':
            data = request.json or {}
            key = data.get('key')
            html = data.get('html', '')
            user_id = data.get('userId') or data.get('user_id')
            if not key:
                return jsonify({'success': False, 'error': 'Signature key required'}), 400
            ok = sync_manager.save_signature_html(key, html, user_id=user_id)
            return jsonify({'success': ok})
        else:
            sigs = sync_manager.load_all_signatures()
            return jsonify({'success': True, 'signatures': sigs})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

# -------------------------------------------------------------
# CONTACT MANAGEMENT & AUTOCOMPLETE ENDPOINTS
# -------------------------------------------------------------
@app.route('/api/contacts', methods=['GET', 'POST'])
def handle_contacts():
    try:
        if request.method == 'POST':
            data = request.json or {}
            name = data.get('name', '')
            email_addr = data.get('email', '')
            company = data.get('company', '')
            notes = data.get('notes', '')
            if not email_addr:
                return jsonify({'success': False, 'error': 'Email is required'}), 400
            contact = sync_manager.upsert_email_contact(name, email_addr, company, notes)
            return jsonify({'success': True, 'contact': contact})
        else:
            q = request.args.get('query', '')
            limit = int(request.args.get('limit', 100))
            contacts = sync_manager.get_email_contacts(query=q, limit=limit)
            return jsonify({'success': True, 'contacts': contacts})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/contacts/bulk', methods=['POST'])
def handle_contacts_bulk():
    try:
        data = request.json or {}
        contacts = data.get('contacts', [])
        count = sync_manager.bulk_upsert_email_contacts(contacts)
        return jsonify({'success': True, 'count': count})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/contacts/<int:contact_id>', methods=['DELETE'])
def handle_contact_delete(contact_id):
    try:
        ok = sync_manager.delete_email_contact(contact_id)
        return jsonify({'success': ok})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/contacts/import', methods=['POST'])
def handle_contacts_import():
    try:
        if 'file' in request.files:
            f = request.files['file']
            file_bytes = f.read()
            fn = f.filename
            imported, total, err = sync_manager.import_contacts_csv_or_excel(file_bytes, fn)
            if err:
                return jsonify({'success': False, 'error': err}), 400
            return jsonify({'success': True, 'imported_count': imported, 'total_rows': total})
        else:
            data = request.json or {}
            contacts = data.get('contacts', [])
            count = sync_manager.bulk_upsert_email_contacts(contacts)
            return jsonify({'success': True, 'imported_count': count})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/contacts/export', methods=['GET'])
def handle_contacts_export():
    try:
        contacts = sync_manager.get_email_contacts(limit=10000)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Name', 'Email', 'Company', 'Notes', 'Frequency', 'Last Used'])
        for c in contacts:
            writer.writerow([c.get('name', ''), c.get('email', ''), c.get('company', ''), c.get('notes', ''), c.get('frequency', 1), c.get('last_used', '')])
        
        output.seek(0)
        bio = io.BytesIO(output.getvalue().encode('utf-8-sig'))
        return send_file(bio, mimetype='text/csv; charset=utf-8', as_attachment=True, download_name='ap_merin_contacts.csv')
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

def native_save_file_dialog(filename, file_type='csv'):
    """
    Prompts the user with the OS native Save File dialog to choose save location.
    Returns chosen file path or None if cancelled.
    """
    import subprocess
    import tempfile
    
    if sys.platform == 'win32':
        filter_str = 'CSV files (*.csv)|*.csv|All files (*.*)|*.*' if file_type == 'csv' else 'All files (*.*)|*.*'
        ps_code = f'''
Add-Type -AssemblyName System.Windows.Forms
$d = New-Object System.Windows.Forms.SaveFileDialog
$d.Title = "Save Contacts CSV"
$d.Filter = "{filter_str}"
$d.FileName = "{filename}"
$d.InitialDirectory = [System.Environment]::GetFolderPath("Desktop")
if ($d.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK) {{
    Write-Output $d.FileName
}} else {{
    Write-Output "__CANCELLED__"
}}
'''
        with tempfile.NamedTemporaryFile('w', suffix='.ps1', delete=False, encoding='utf-8') as tf:
            tf.write(ps_code)
            ps_path = tf.name
        try:
            res = subprocess.run(
                ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-STA', '-File', ps_path],
                capture_output=True, text=True, timeout=120
            )
            out = res.stdout.strip().splitlines()
            result = out[-1].strip() if out else '__CANCELLED__'
            if result and result != '__CANCELLED__':
                return result
            return None
        except Exception as e:
            print(f"Windows save dialog error: {e}")
            return None
        finally:
            if os.path.exists(ps_path):
                try:
                    os.remove(ps_path)
                except Exception:
                    pass

    elif sys.platform == 'darwin':
        script = f'''
set defaultPath to (path to desktop folder as text)
try
    set chosenFile to choose file name with prompt "Save Contacts CSV" default name "{filename}" default location defaultPath
    return POSIX path of chosenFile
on error
    return "__CANCELLED__"
end try
'''
        try:
            res = subprocess.run(['osascript', '-e', script], capture_output=True, text=True, timeout=120)
            result = res.stdout.strip()
            if result and result != '__CANCELLED__':
                return result
            return None
        except Exception as e:
            print(f"macOS save dialog error: {e}")
            return None

    return None

@app.route('/api/save-file-dialog', methods=['POST'])
def handle_save_file_dialog():
    try:
        data = request.json or {}
        filename = data.get('filename', 'export.csv')
        content = data.get('content', '')
        file_type = data.get('file_type', 'csv')
        
        target_path = native_save_file_dialog(filename, file_type=file_type)
        if target_path:
            with open(target_path, 'w', encoding='utf-8-sig') as f:
                f.write(content)
            return jsonify({'success': True, 'path': target_path})
        else:
            return jsonify({'success': False, 'cancelled': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/status', methods=['GET'])
def handle_sync_status():
    try:
        status = sync_manager.get_sync_status()
        return jsonify(status)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/configure', methods=['POST'])
def handle_sync_configure():
    try:
        data = request.json or {}
        folder = data.get('shared_folder', '')
        res = sync_manager.configure_shared_folder(folder)
        return jsonify(res)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/cloud-folders', methods=['GET'])
def handle_sync_cloud_folders():
    try:
        folders = sync_manager.detect_cloud_folders()
        return jsonify({'success': True, 'folders': folders})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/comparisons', methods=['GET'])
def handle_sync_list_comparisons():
    try:
        limit = int(request.args.get('limit', 50))
        comps = sync_manager.list_comparisons(limit=limit)
        return jsonify({'success': True, 'comparisons': comps})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/comparisons/save', methods=['POST'])
def handle_sync_save_comparison():
    try:
        data = request.json or {}
        res = sync_manager.save_comparison(data)
        return jsonify(res)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/comparisons/<comp_id>', methods=['GET'])
def handle_sync_get_comparison(comp_id):
    try:
        comp = sync_manager.get_comparison(comp_id)
        if not comp:
            return jsonify({'success': False, 'error': 'Comparison not found'}), 404
        return jsonify({'success': True, 'comparison': comp})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/comparisons/<comp_id>', methods=['DELETE'])
def handle_sync_delete_comparison(comp_id):
    try:
        res = sync_manager.delete_comparison(comp_id)
        return jsonify(res)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/export', methods=['GET'])
def handle_sync_export():
    try:
        bundle = sync_manager.export_sync_bundle()
        raw_json = json.dumps(bundle, indent=2)
        out_stream = io.BytesIO(raw_json.encode('utf-8'))
        out_stream.seek(0)
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        return send_file(
            out_stream,
            mimetype="application/json",
            as_attachment=True,
            download_name=f"AP_Merin_Sync_{timestamp}.apmerin_sync"
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/sync/import', methods=['POST'])
def handle_sync_import():
    try:
        if 'file' in request.files:
            file_obj = request.files['file']
            bundle_dict = json.loads(file_obj.read().decode('utf-8'))
        else:
            bundle_dict = request.json or {}

        if not bundle_dict:
            return jsonify({'success': False, 'error': 'No sync data received'}), 400

        res = sync_manager.import_sync_bundle(bundle_dict)
        return jsonify(res)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/download-file', methods=['POST'])
def handle_download_file():
    try:
        data = request.get_json(silent=True) or {}
        b64_str = data.get('base64', '')
        filename = data.get('filename', 'download.xlsx')
        mime_type = data.get('mime_type', 'application/octet-stream')
        if not b64_str:
            return jsonify({'success': False, 'error': 'No content provided'}), 400

        clean_b64 = re.sub(r'^data:.*?;base64,', '', b64_str).strip()
        clean_b64 = re.sub(r'[\r\n\s]', '', clean_b64)
        raw_bytes = base64.b64decode(clean_b64)
        out_stream = io.BytesIO(raw_bytes)
        out_stream.seek(0)

        return send_file(
            out_stream,
            mimetype=mime_type,
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


if __name__ == '__main__':
    print("Starting AP Merin Quotation App with Universal Document Previews on http://127.0.0.1:5000")
    app.run(host='0.0.0.0', port=5000, debug=True)


